"""
El documento de una plantilla en Word (DOCX) editable.

Recorre el mismo ``RenderedDocument`` que el PDF, así que el contenido es idéntico. Usa los estilos
integrados de Word (Título, Título 1, Normal, Viñeta de lista): quien lo abra puede cambiar el
formato de todo el documento desde Word o LibreOffice sin tocar párrafo a párrafo.
"""

import io
import logging
import re
from datetime import datetime
from typing import Optional, Tuple

from docx import Document
from docx.shared import Inches

from .render import KIND_BULLETS, RenderedDocument

logger = logging.getLogger(__name__)

_BOLD = re.compile(r"\*\*(?!\s)(.+?)(?<!\s)\*\*", re.DOTALL)

_NOTE = (
    "Borrador generado por Ellysia a partir de los datos de la empresa. La dirección debe "
    "revisarlo, adaptarlo a la realidad de la organización y aprobarlo antes de darlo por vigente. "
    "No constituye asesoramiento jurídico."
)


def _add_runs(paragraph, text: str) -> None:
    """Escribe un texto en un párrafo: ``**negrita**`` y los saltos de línea de verdad."""
    position = 0
    for match in _BOLD.finditer(text):
        _add_plain(paragraph, text[position:match.start()])
        _add_plain(paragraph, match.group(1), bold=True)
        position = match.end()
    _add_plain(paragraph, text[position:])


def _add_plain(paragraph, text: str, bold: bool = False) -> None:
    """Añade texto sin marcas, convirtiendo cada salto de línea en un salto de Word."""
    for number, line in enumerate(text.split("\n")):
        if number:
            paragraph.add_run().add_break()
        if line:
            paragraph.add_run(line).bold = bold or None


def build_docx(
    *,
    document: RenderedDocument,
    company_name: str,
    tax_id: str = "",
    framework_label: str = "",
    logo: Optional[Tuple[bytes, str]] = None,
    author: str = "Ellysia Security Team",
    generated_at: Optional[datetime] = None,
) -> bytes:
    """Compone el documento en formato Word.

    Args:
        document: El documento ya sin marcadores (``render_document``).
        company_name: Razón social para la portada.
        tax_id: NIF de la empresa. Por defecto vacío.
        framework_label: Nombre corto del marco de la plantilla. Por defecto vacío.
        logo: Logotipo de la empresa como ``(bytes, extensión)`` (``decode_logo``). Por defecto
            ``None``; uno que no se pueda insertar se omite.
        author: Autor de las propiedades del documento.
        generated_at: Fecha de la portada. Por defecto ahora.

    Returns:
        bytes: El fichero ``.docx``.
    """
    moment = generated_at or datetime.now()
    word = Document()
    word.core_properties.title = document.title
    word.core_properties.author = author

    if logo is not None:
        try:
            word.add_picture(io.BytesIO(logo[0]), width=Inches(1.6))
        except Exception:  # noqa: BLE001 — un logotipo ilegible no debe impedir el documento
            logger.warning("El logotipo de la empresa no se pudo insertar; se genera sin él", exc_info=True)

    word.add_paragraph(document.title, style="Title")
    if company_name:
        word.add_paragraph(company_name, style="Subtitle")
    for label, value in (("Empresa", company_name), ("NIF", tax_id), ("Marco", framework_label),
                         ("Fecha", moment.strftime("%d/%m/%Y"))):
        if value:
            paragraph = word.add_paragraph()
            paragraph.add_run(f"{label}: ").bold = True
            paragraph.add_run(value)

    for section in document.sections:
        word.add_heading(section.heading, level=1)
        for block in section.blocks:
            if block.kind == KIND_BULLETS:
                for text in block.texts:
                    _add_runs(word.add_paragraph(style="List Bullet"), text)
            else:
                for text in block.texts:
                    _add_runs(word.add_paragraph(style="Normal"), text)

    note = word.add_paragraph()
    note.add_run(_NOTE).italic = True

    buffer = io.BytesIO()
    word.save(buffer)
    return buffer.getvalue()
