"""Los documentos de plantilla: el contenido sale igual para PDF y Word, y el PDF es válido."""

import base64
import re
from datetime import datetime

import pytest

from src.modules.features.eunomia.services.documents import missing_required, render_document
from src.modules.features.eunomia.services.documents.pdf import EunomiaTemplatePDF, decode_logo
from src.modules.features.eunomia.services.templates import get_template, load_templates

pytestmark = pytest.mark.unit

def _png() -> str:
    """Un PNG de 8x8 píxeles, en base64."""
    import io

    from PIL import Image

    buffer = io.BytesIO()
    Image.new("RGB", (8, 8), "#3f5a8c").save(buffer, format="PNG")
    return base64.b64encode(buffer.getvalue()).decode()


def _values(template, **overrides):
    values = {}
    for field in template.fields:
        values[field.key] = {
            "date": "2026-10-10", "list": "Uno\n- Dos\nTres",
        }.get(field.type, f"valor de {field.key}")
    values.update(overrides)
    return values


def _pdf(template, values, **options):
    document = render_document(template, values)
    return EunomiaTemplatePDF(
        document=document, company_name=values.get("company_name", ""), tax_id="B12345674",
        framework_label="NIS2", generated_at=datetime(2026, 10, 10), **options,
    ).generate()


@pytest.mark.parametrize("template", load_templates(), ids=lambda t: t.key)
def test_every_template_generates_a_valid_pdf(template):
    pdf = _pdf(template, _values(template))

    assert pdf.startswith(b"%PDF") and len(pdf) > 2000


def test_a_value_with_markup_characters_does_not_break_the_document():
    template = get_template("security-policy")
    values = _values(template, company_name="Fish & <Chips> S.L.", scope="a < b && c > d <script>")

    assert _pdf(template, values).startswith(b"%PDF")


def test_the_company_logo_is_embedded_when_there_is_one():
    template = get_template("security-policy")
    values = _values(template)

    without = _pdf(template, values)
    with_logo = _pdf(template, values, logo=decode_logo(f"data:image/png;base64,{_png()}"))

    assert with_logo != without
    assert b"/Subtype /Image" in with_logo


def test_without_a_company_logo_the_cover_carries_the_eunomia_one():
    template = get_template("security-policy")

    assert b"/Subtype /Image" in _pdf(template, _values(template))


@pytest.mark.parametrize("uri", [None, "", "data:text/html;base64,PGI+", "data:image/png;base64,@@@", "no-uri"])
def test_an_unusable_logo_is_ignored(uri):
    assert decode_logo(uri) is None


def test_a_list_alone_in_its_paragraph_becomes_bullets_and_the_date_is_formatted():
    template = get_template("security-policy")

    document = render_document(template, _values(template))

    objectives = next(s for s in document.sections if s.heading.startswith("3."))
    bullets = next(b for b in objectives.blocks if b.kind == "bullets")
    assert bullets.texts == ("Uno", "Dos", "Tres")
    approval = next(s for s in document.sections if s.heading.startswith("6."))
    assert "10/10/2026" in approval.blocks[0].texts[0]


def test_an_optional_empty_field_prints_a_dash_and_no_marker_survives():
    template = get_template("incident-procedure")

    document = render_document(template, _values(template, own_criteria=""))

    text = " ".join(t for s in document.sections for b in s.blocks for t in b.texts)
    assert "—" in text
    assert not re.search(r"\{\{|\}\}", text)


def test_missing_required_lists_what_blocks_generation():
    template = get_template("incident-procedure")

    assert missing_required(template, {"company_name": "Acme"}) >= ["tax_id"]
    assert missing_required(template, _values(template)) == []


def test_a_corrupt_logo_does_not_prevent_the_document():
    template = get_template("security-policy")
    corrupt = decode_logo("data:image/png;base64," + base64.b64encode(b"no soy un png").decode())

    assert _pdf(template, _values(template), logo=corrupt).startswith(b"%PDF")


# ── Word ───────────────────────────────────────────────────────────────────

def _docx(template, values, **options):
    import io

    from docx import Document

    from src.modules.features.eunomia.services.documents.docx import build_docx

    data = build_docx(document=render_document(template, values), company_name=values.get("company_name", ""),
                      tax_id="B12345674", framework_label="NIS2", generated_at=datetime(2026, 10, 10), **options)
    return Document(io.BytesIO(data))


@pytest.mark.parametrize("template", load_templates(), ids=lambda t: t.key)
def test_every_template_generates_a_docx_with_all_its_sections(template):
    word = _docx(template, _values(template))

    headings = [p.text for p in word.paragraphs if p.style.name == "Heading 1"]
    assert headings == [section.heading for section in render_document(template, _values(template)).sections]


def test_the_docx_uses_the_built_in_word_styles():
    template = get_template("security-policy")

    word = _docx(template, _values(template))

    used = {p.style.name for p in word.paragraphs}
    assert used <= {"Title", "Subtitle", "Heading 1", "Normal", "List Bullet"}
    assert {"Title", "Heading 1", "List Bullet"} <= used


def test_the_docx_carries_the_same_text_as_the_rendered_document():
    template = get_template("incident-procedure")
    values = _values(template, company_name="Acme & <Co>")

    word = _docx(template, values)

    text = "\n".join(p.text for p in word.paragraphs)
    assert "Acme & <Co>" in text and "24 horas" in text
    assert "**" not in text


def test_the_docx_embeds_the_logo_and_survives_a_corrupt_one():
    template = get_template("security-policy")
    good = decode_logo(f"data:image/png;base64,{_png()}")
    corrupt = decode_logo("data:image/png;base64," + base64.b64encode(b"no soy un png").decode())

    assert len(_docx(template, _values(template), logo=good).inline_shapes) == 1
    assert len(_docx(template, _values(template), logo=corrupt).inline_shapes) == 0
